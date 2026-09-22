from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from tools.build_catalog import build_catalog


def test_build_catalog_preserves_source_and_fills_navigation_gaps(
    tmp_path: Path,
):
    source = tmp_path / "catalog.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["No", "Schema", "Modalità", "Attività", "Scala", "Livello", "Descrittore"])
    sheet.append(
        [
            17,
            "Competenze nella lingua dei segni",
            None,
            "Pragmatica",
            "Strutturazione del testo",
            "B1+",
            "Descrittore di prova",
        ]
    )
    workbook.save(source)

    rows, report = build_catalog(
        source,
        tmp_path / "missing-sample.json",
        expected_rows=1,
    )

    assert report["catalog_rows"] == 1
    assert report["blank_source_modality"] == 1
    assert rows[0]["modality"] == "Pragmatica"
    assert rows[0]["activity"] == "Scale generali"
    assert rows[0]["source_modality"] == ""
    assert rows[0]["correct_level"] == "B1+"


def test_build_catalog_keeps_src_8_editorial_correction_on_rebuild(
    tmp_path: Path,
):
    source = tmp_path / "catalog.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["No", "Schema", "Modalità", "Attività", "Scala", "Livello", "Descrittore"])
    sheet.append(
        [
            8,
            "Attività linguistico-comunicative",
            "Ricezione",
            "Comprensione orale",
            "Comprensione orale generale",
            "B1+",
            "È in grado di comprendere un discorso in un varietà piuttosto familiare.",
        ]
    )
    workbook.save(source)

    rows, _ = build_catalog(source, tmp_path / "missing-sample.json", expected_rows=1)

    assert "in una varietà" in rows[0]["descriptor_text"]
    assert "in un varietà" in load_workbook(source).active["G2"].value


def test_build_catalog_keeps_src_11_signing_adaptation_on_rebuild(
    tmp_path: Path,
):
    source = tmp_path / "catalog.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["No", "Schema", "Modalità", "Attività", "Scala", "Livello", "Descrittore"])
    sheet.append(
        [
            11,
            "Attività linguistico-comunicative",
            "Ricezione",
            "Comprensione orale",
            "Comprensione orale generale",
            "A2",
            "È in grado di comprendere espressioni, purché si parli lentamente e chiaramente.",
        ]
    )
    workbook.save(source)

    rows, _ = build_catalog(source, tmp_path / "missing-sample.json", expected_rows=1)

    assert "si parli/segni lentamente e chiaramente" in rows[0]["descriptor_text"]
    assert "si parli lentamente e chiaramente" in load_workbook(source).active["G2"].value


def test_build_catalog_keeps_src_13_signing_adaptation_on_rebuild(
    tmp_path: Path,
):
    source = tmp_path / "catalog.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["No", "Schema", "Modalità", "Attività", "Scala", "Livello", "Descrittore"])
    sheet.append(
        [
            13,
            "Attività linguistico-comunicative",
            "Ricezione",
            "Comprensione orale",
            "Comprensione orale generale",
            "A1",
            "È in grado di cogliere un dato, purché si parli lentamente e chiaramente.",
        ]
    )
    workbook.save(source)

    rows, _ = build_catalog(source, tmp_path / "missing-sample.json", expected_rows=1)

    assert "si parli/segni lentamente e chiaramente" in rows[0]["descriptor_text"]
    assert "si parli lentamente e chiaramente" in load_workbook(source).active["G2"].value
